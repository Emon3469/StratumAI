from pathlib import Path


def _parse_docx(file_path: str) -> str:
    from docx import Document

    document = Document(file_path)
    blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return "\n\n".join(blocks)


def _parse_pptx(file_path: str) -> str:
    from pptx import Presentation

    presentation = Presentation(file_path)
    slides: list[str] = []
    for number, slide in enumerate(presentation.slides, start=1):
        texts = [f"Slide {number}"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
                if text.strip():
                    texts.append(text.strip())
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                texts.append(f"Notes: {notes}")
        slides.append("\n".join(texts))
    return "\n\n".join(slides)


def parse_office(file_path: str) -> str:
    suffix = Path(file_path).suffix.lower()
    if suffix == ".docx":
        return _parse_docx(file_path)
    if suffix == ".pptx":
        return _parse_pptx(file_path)
    raise ValueError(f"Unsupported office format: {suffix}")
